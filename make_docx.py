import os
import sys
from PIL import Image as PILImage

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

import report_content as rc

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_header_footer(doc):
    for s in doc.sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)
        s.different_first_page_header_footer = True
        
        # Header (Pages > 1)
        header = s.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("TripMeta — Intelligent Travel Recommendation Platform | SPIT | AY 2026–27")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(100, 116, 139)
        
        # Footer (Pages > 1)
        footer = s.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        frun1 = fp.add_run("Department of Computer Engineering, S.P.I.T., Mumbai\t\t")
        frun1.font.name = "Calibri"
        frun1.font.size = Pt(8.5)
        frun1.font.color.rgb = RGBColor(100, 116, 139)

def format_table(doc_table, matrix, col_widths=None):
    doc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    doc_table.autofit = False
    
    # Header row
    hdr_cells = doc_table.rows[0].cells
    for c_idx, val in enumerate(matrix[0]):
        cell = hdr_cells[c_idx]
        cell.text = str(val)
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(8.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            
    # Data rows
    for r_idx, row_data in enumerate(matrix[1:]):
        row_cells = doc_table.rows[r_idx + 1].cells
        bg_color = "F8FAFC" if (r_idx % 2 == 1) else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = row_cells[c_idx]
            cell.text = str(val)
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = "Calibri"
                run.font.size = Pt(8)
                run.font.color.rgb = RGBColor(30, 41, 59)
                
    # Column widths
    if col_widths:
        for row in doc_table.rows:
            for c_idx, w in enumerate(col_widths):
                if c_idx < len(row.cells):
                    row.cells[c_idx].width = Inches(w)

def add_image_safe(doc, path, caption, max_width=6.0):
    if not os.path.exists(path):
        p = doc.add_paragraph(f"[Image file not found: {path}]")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return
    try:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(path, width=Inches(max_width))
        
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        crun = cp.add_run(caption)
        crun.font.name = "Calibri"
        crun.font.size = Pt(8.5)
        crun.font.bold = True
        crun.font.color.rgb = RGBColor(71, 85, 105)
    except Exception as e:
        p = doc.add_paragraph(f"[Error embedding image {path}: {str(e)}]")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

def create_docx(output_filename="TripMeta_Final_Project_Report.docx"):
    print(f"Building DOCX report: {output_filename}")
    doc = docx.Document()
    add_header_footer(doc)

    # 1. COVER PAGE
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("A PROJECT REPORT ON")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(10)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(100, 116, 139)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run(rc.METADATA["title"].upper())
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(30, 58, 138) # #1E3A8A

    p_deg = doc.add_paragraph()
    p_deg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_deg1 = p_deg.add_run("Submitted in partial fulfillment of the requirements for the degree of\n")
    r_deg1.font.name = "Calibri"
    r_deg1.font.size = Pt(11)
    r_deg1.font.color.rgb = RGBColor(71, 85, 105)
    r_deg2 = p_deg.add_run(rc.METADATA["degree"])
    r_deg2.font.name = "Calibri"
    r_deg2.font.size = Pt(12)
    r_deg2.font.bold = True
    r_deg2.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph() # spacer

    # Authors Table
    t_auth = doc.add_table(rows=2, cols=2)
    t_auth.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_auth.rows[0].cells[0].paragraphs[0].add_run("Submitted By:").bold = True
    t_auth.rows[0].cells[1].paragraphs[0].add_run("Under the Guidance of:").bold = True
    for cell in t_auth.rows[0].cells:
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(30, 58, 138)
            run.font.size = Pt(10)

    p0 = t_auth.rows[1].cells[0].paragraphs[0]
    p0.add_run(f"{rc.METADATA['students'][0]['name']} (UID: {rc.METADATA['students'][0]['uid']})\n")
    p0.add_run(f"{rc.METADATA['students'][1]['name']} (UID: {rc.METADATA['students'][1]['uid']})")

    p1 = t_auth.rows[1].cells[1].paragraphs[0]
    p1.add_run(f"{rc.METADATA['guide']['name']}\n{rc.METADATA['guide']['designation']}\nDepartment of Computer Engineering")

    doc.add_paragraph() # spacer
    doc.add_paragraph()

    # Institute block
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_dept = p_inst.add_run(f"{rc.METADATA['department']}\n")
    r_dept.font.bold = True
    r_dept.font.size = Pt(12)
    r_dept.font.color.rgb = RGBColor(30, 58, 138)

    r_col = p_inst.add_run(f"{rc.METADATA['institution']}\n{rc.METADATA['location']}\n")
    r_col.font.bold = True
    r_col.font.size = Pt(11)
    r_col.font.color.rgb = RGBColor(15, 23, 42)

    r_ay = p_inst.add_run(f"Academic Year {rc.METADATA['academic_year']}")
    r_ay.font.bold = True
    r_ay.font.size = Pt(11)
    r_ay.font.color.rgb = RGBColor(37, 99, 235)

    doc.add_page_break()

    # 2. CERTIFICATE
    h_cert = doc.add_heading("CERTIFICATE", level=1)
    h_cert.runs[0].font.color.rgb = RGBColor(30, 58, 138)
    for p in rc.CERTIFICATE_TEXT.split('\n\n'):
        doc.add_paragraph(p)
    doc.add_paragraph()
    
    t_cert = doc.add_table(rows=2, cols=2)
    t_cert.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_cert.rows[0].cells[0].paragraphs[0].add_run(f"{rc.METADATA['guide']['name']}\nProject Guide\nAssistant Professor")
    t_cert.rows[0].cells[1].paragraphs[0].add_run(f"{rc.METADATA['hod']['name']}\nHead of Department\nComputer Engineering")
    t_cert.rows[1].cells[0].paragraphs[0].add_run(f"\n\n{rc.METADATA['principal']['name']}\nPrincipal\nS.P.I.T., Mumbai")
    t_cert.rows[1].cells[1].paragraphs[0].add_run("\n\nExternal Examiner\nSignature & Seal")
    doc.add_page_break()

    # 3. DECLARATION
    h_decl = doc.add_heading("DECLARATION", level=1)
    h_decl.runs[0].font.color.rgb = RGBColor(30, 58, 138)
    for p in rc.DECLARATION_TEXT.split('\n\n'):
        doc.add_paragraph(p)
    doc.add_paragraph()
    doc.add_paragraph(f"{rc.METADATA['students'][0]['name']} (UID: {rc.METADATA['students'][0]['uid']})\n{rc.METADATA['students'][1]['name']} (UID: {rc.METADATA['students'][1]['uid']})\nPlace: Mumbai\nDate: October 30, 2026")
    doc.add_page_break()

    # 4. ACKNOWLEDGEMENT
    h_ack = doc.add_heading("ACKNOWLEDGEMENT", level=1)
    h_ack.runs[0].font.color.rgb = RGBColor(30, 58, 138)
    for p in rc.ACKNOWLEDGEMENT_TEXT.split('\n\n'):
        doc.add_paragraph(p)
    doc.add_paragraph()
    doc.add_paragraph(f"{rc.METADATA['students'][0]['name']} & {rc.METADATA['students'][1]['name']}")
    doc.add_page_break()

    # 5. ABSTRACT
    h_abs = doc.add_heading("ABSTRACT", level=1)
    h_abs.runs[0].font.color.rgb = RGBColor(30, 58, 138)
    for p in rc.ABSTRACT_TEXT.split('\n\n'):
        doc.add_paragraph(p)
    doc.add_paragraph()
    p_kw = doc.add_paragraph()
    r_kwh = p_kw.add_run("Keywords: ")
    r_kwh.bold = True
    p_kw.add_run("Tourism Recommender System, Vector Space Model, TF-IDF Vectorization, Cosine Similarity, Climatological Seasonal Weighting, Constraint Penalties, Itinerary Planning, Single Page Application, React 19, SQLite.")
    doc.add_page_break()

    # 5.5 TABLE OF CONTENTS
    h_toc = doc.add_heading("TABLE OF CONTENTS", level=1)
    h_toc.runs[0].font.color.rgb = RGBColor(30, 58, 138)
    toc_mat = [["Chapter / Section Title", "Page"]]
    import make_pdf
    # Use the captured page map if available, or standard targets
    ch1_p = make_pdf.GLOBAL_PAGE_MAP.get("Chapter 1: Introduction and Problem Definition", 10)
    def docx_pg(key, fallback):
        raw = make_pdf.GLOBAL_PAGE_MAP.get(key)
        if raw is None:
            return fallback
        if raw < ch1_p:
            return make_pdf.to_roman(raw)
        return str(raw - ch1_p + 1)

    toc_raw = [
        ("Certificate", docx_pg("Certificate", "ii")),
        ("Declaration", docx_pg("Declaration", "iii")),
        ("Acknowledgement", docx_pg("Acknowledgement", "iv")),
        ("Abstract", docx_pg("Abstract", "v")),
        ("List of Abbreviations", docx_pg("List of Abbreviations", "viii")),
        ("List of Tables", docx_pg("List of Tables", "ix")),
        ("List of Figures", docx_pg("List of Figures", "ix")),
        ("Chapter 1: Introduction and Problem Definition", docx_pg("Chapter 1: Introduction and Problem Definition", "1")),
        ("   1.1 Background and Context", docx_pg("1.1 Background and Context", "1")),
        ("   1.2 Problem Statement", docx_pg("1.2 Problem Statement", "1")),
        ("   1.3 Motivation", docx_pg("1.3 Motivation", "2")),
        ("   1.4 Objectives of the Project", docx_pg("1.4 Objectives of the Project", "2")),
        ("   1.5 Scope of the Project", docx_pg("1.5 Scope of the Project", "2")),
        ("   1.6 Proposed Solution Overview", docx_pg("1.6 Proposed Solution Overview", "3")),
        ("   1.7 Novelty and Innovation", docx_pg("1.7 Novelty and Innovation", "3")),
        ("   1.8 Feasibility Analysis", docx_pg("1.8 Feasibility Analysis", "3")),
        ("   1.9 Social and Environmental Relevance", docx_pg("1.9 Social and Environmental Relevance", "4")),
        ("   1.10 Key Performance Indicators (KPIs)", docx_pg("1.10 Key Performance Indicators (KPIs)", "4")),
        ("Chapter 2: Literature Review and Study of Existing Systems", docx_pg("Chapter 2: Literature Review and Study of Existing Systems", "5")),
        ("   2.1 Overview of Modern Recommender Systems in Tourism", docx_pg("2.1 Overview of Modern Recommender Systems in Tourism", "5")),
        ("   2.2 Comparative Study of Commercial Platforms", docx_pg("2.2 Comparative Study of Commercial Platforms", "5")),
        ("   2.3 Critical Limitations of Existing Systems", docx_pg("2.3 Critical Limitations of Existing Systems", "6")),
        ("   2.4 Research and Engineering Gap", docx_pg("2.4 Research and Engineering Gap", "6")),
        ("   2.5 Justification of the Selected Algorithmic Approach", docx_pg("2.5 Justification of the Selected Algorithmic Approach", "7")),
        ("Chapter 3: Software Requirements Specification (SRS)", docx_pg("Chapter 3: Software Requirements Specification (SRS)", "8")),
        ("   3.1 User Personas and Stakeholder Identification", docx_pg("3.1 User Personas and Stakeholder Identification", "8")),
        ("   3.2 Functional Requirements (FR-01 to FR-14)", docx_pg("3.2 Functional Requirements", "8")),
        ("   3.3 Non-Functional Requirements", docx_pg("3.3 Non-Functional Requirements", "9")),
        ("   3.4 Hardware and Infrastructure Requirements", docx_pg("3.4 Hardware and Infrastructure Requirements", "10")),
        ("   3.5 Software and Technology Stack Specifications", docx_pg("3.5 Software and Technology Stack Specifications", "10")),
        ("   3.6 Detailed Use Case Specifications", docx_pg("3.6 Detailed Use Case Specifications", "11")),
        ("Chapter 4: System Architecture and Design", docx_pg("Chapter 4: System Architecture and Design", "12")),
        ("   4.1 High-Level Layered Architecture", docx_pg("4.1 High-Level Layered Architecture", "12")),
        ("   4.2 UML Use Case Diagram", docx_pg("4.2 UML Use Case Diagram", "13")),
        ("   4.3 UML Sequence Diagram: Recommendation Flow", docx_pg("4.3 UML Sequence Diagram: Recommendation Flow", "14")),
        ("   4.4 Data Flow Diagrams (DFD)", docx_pg("4.4 Data Flow Diagrams (DFD)", "14")),
        ("   4.5 Relational Database Schema and ER Diagram", docx_pg("4.5 Relational Database Schema and ER Diagram", "15")),
        ("   4.6 Data Dictionary and Table Structures", docx_pg("4.6 Data Dictionary and Table Structures", "15")),
        ("   4.7 RESTful API Specifications", docx_pg("4.7 RESTful API Specifications", "16")),
        ("Chapter 5: Algorithm Design and Data Structures", docx_pg("Chapter 5: Algorithm Design and Data Structures", "17")),
        ("   5.1 Mathematical Foundations of Vector Space Modeling", docx_pg("5.1 Mathematical Foundations of Vector Space Modeling (VSM)", "17")),
        ("   5.2 Term Frequency (TF) & Inverse Document Frequency (IDF)", docx_pg("5.2 Term Frequency (TF) and Inverse Document Frequency (IDF) Formulation", "17")),
        ("   5.3 Vector Normalization and Cosine Similarity Formulation", docx_pg("5.3 Vector Normalization and Cosine Similarity Formulation", "18")),
        ("   5.4 Hybrid Multi-Criteria Composite Scoring Model", docx_pg("5.4 Hybrid Multi-Criteria Composite Scoring Model", "18")),
        ("   5.5 Heuristic Penalty Constraint Modeling", docx_pg("5.5 Heuristic Penalty Constraint Modeling", "19")),
        ("   5.6 Algorithmic Pipeline and Formal Pseudocode", docx_pg("5.6 Algorithmic Pipeline and Formal Pseudocode", "19")),
        ("   5.7 Time and Space Complexity Analysis", docx_pg("5.7 Time and Space Complexity Analysis", "20")),
        ("Chapter 6: Subject-Wise Engineering Concepts & APNI Leap Alignment", docx_pg("Chapter 6: Subject-Wise Engineering Concepts & APNI Leap Alignment", "21")),
        ("   6.1 Software Engineering (SE)", docx_pg("6.1 Software Engineering (SE)", "21")),
        ("   6.2 Computer Networks and Web Technologies (CN & WT)", docx_pg("6.2 Computer Networks and Web Technologies (CN & WT)", "21")),
        ("   6.3 Database Management Systems (DBMS)", docx_pg("6.3 Database Management Systems (DBMS)", "22")),
        ("   6.4 Operating Systems and Concurrency (OS)", docx_pg("6.4 Operating Systems and Concurrency (OS)", "22")),
        ("   6.5 Data Structures and Algorithms (DSA & DAA)", docx_pg("6.5 Data Structures and Design & Analysis of Algorithms (DSA & DAA)", "23")),
        ("Chapter 7: Implementation and Module-Wise Coding", docx_pg("Chapter 7: Implementation and Module-Wise Coding", "24")),
        ("   7.1 Development Environment and Codebase Structure", docx_pg("7.1 Development Environment and Codebase Structure", "24")),
        ("   7.2 Frontend Implementation Details", docx_pg("7.2 Frontend Implementation Details", "24")),
        ("   7.3 Backend Implementation Details", docx_pg("7.3 Backend Implementation Details", "25")),
        ("   7.4 Catalog Seeding and Image Synchronization Subsystem", docx_pg("7.4 Catalog Seeding and Image Synchronization Subsystem", "25")),
        ("   7.5 Team Member Contributions and Work Allocation", docx_pg("7.5 Team Member Contributions and Work Allocation", "26")),
        ("   7.6 System Integration Challenges and Technical Solutions", docx_pg("7.6 System Integration Challenges and Technical Solutions", "26")),
        ("Chapter 8: Scalability, Performance and Resource Management", docx_pg("Chapter 8: Scalability, Performance and Resource Management", "27")),
        ("   8.1 Workload Modeling and Concurrency Profile", docx_pg("8.1 Workload Modeling and Concurrency Profile", "27")),
        ("   8.2 Computational Bottleneck Identification", docx_pg("8.2 Computational Bottleneck Identification", "27")),
        ("   8.3 Performance Optimization Techniques", docx_pg("8.3 Performance Optimization Techniques", "27")),
        ("   8.4 Empirical Benchmark Results", docx_pg("8.4 Empirical Benchmark Results", "28")),
        ("   8.5 Resource Utilization and Memory Footprint", docx_pg("8.5 Resource Utilization and Memory Footprint", "28")),
        ("Chapter 9: Testing and Validation", docx_pg("Chapter 9: Testing and Validation", "29")),
        ("   9.1 Testing Philosophy and Methodology", docx_pg("9.1 Testing Philosophy and Methodology", "29")),
        ("   9.2 Test Suite Architecture and Distribution", docx_pg("9.2 Test Suite Architecture and Distribution", "29")),
        ("   9.3 Comprehensive Test Case Matrix (TC-01 to TC-14)", docx_pg("9.3 Comprehensive Test Case Matrix", "30")),
        ("   9.4 Validation Against Project Objectives and KPIs", docx_pg("9.4 Validation Against Project Objectives and KPIs", "31")),
        ("Chapter 10: Results, Demonstration and Discussion", docx_pg("Chapter 10: Results, Demonstration and Discussion", "32")),
        ("   10.1 System Walkthrough and Visual Demonstrations", docx_pg("10.1 System Walkthrough and Visual Demonstrations", "32")),
        ("   10.2 Experimental Evaluation and User Experience Feedback", docx_pg("10.2 Experimental Evaluation and User Experience Feedback", "34")),
        ("   10.3 Discussion of Results and Comparative Strengths", docx_pg("10.3 Discussion of Results and Comparative Strengths", "34")),
        ("Chapter 11: Conclusion and Future Scope", docx_pg("Chapter 11: Conclusion and Future Scope", "35")),
        ("   11.1 Conclusion and Summary of Work", docx_pg("11.1 Conclusion and Summary of Work", "35")),
        ("   11.2 Key Engineering Contributions", docx_pg("11.2 Key Engineering Contributions", "35")),
        ("   11.3 System Limitations", docx_pg("11.3 System Limitations", "35")),
        ("   11.4 Future Scope and Roadmap", docx_pg("11.4 Future Scope and Roadmap", "36")),
        ("   11.5 Key Learning Outcomes and Reflections", docx_pg("11.5 Key Learning Outcomes and Reflections", "36")),
        ("References", docx_pg("References", "37")),
        ("Appendix A: Academic Integrity & Verification Statement", docx_pg("Appendix A", "38"))
    ]
    for tit, pg in toc_raw:
        toc_mat.append([tit, str(pg)])
    t_toc = doc.add_table(rows=len(toc_mat), cols=2)
    format_table(t_toc, toc_mat, col_widths=[5.6, 1.4])
    doc.add_page_break()

    # 6. LIST OF ABBREVIATIONS
    h_ab = doc.add_heading("LIST OF ABBREVIATIONS", level=1)
    h_ab.runs[0].font.color.rgb = RGBColor(30, 58, 138)
    abbr_mat = [["Abbreviation", "Full Form / Definition"]]
    for ab, d in rc.ABBREVIATIONS:
        abbr_mat.append([ab, d])
    t_ab = doc.add_table(rows=len(abbr_mat), cols=2)
    format_table(t_ab, abbr_mat, col_widths=[1.8, 5.2])
    doc.add_page_break()

    # 7. LIST OF TABLES & LIST OF FIGURES
    h_lt = doc.add_heading("LIST OF TABLES", level=1)
    h_lt.runs[0].font.color.rgb = RGBColor(30, 58, 138)
    tbl_mat = [["Table Number", "Table Title"]]
    for num, tit in rc.TABLES_LIST:
        tbl_mat.append([num, tit])
    t_lt = doc.add_table(rows=len(tbl_mat), cols=2)
    format_table(t_lt, tbl_mat, col_widths=[1.5, 5.5])
    doc.add_paragraph()

    h_lf = doc.add_heading("LIST OF FIGURES", level=1)
    h_lf.runs[0].font.color.rgb = RGBColor(30, 58, 138)
    fig_mat = [["Figure Number", "Figure Title"]]
    for num, tit in rc.FIGURES_LIST:
        fig_mat.append([num, tit])
    t_lf = doc.add_table(rows=len(fig_mat), cols=2)
    format_table(t_lf, fig_mat, col_widths=[1.5, 5.5])
    doc.add_page_break()

    # CHAPTER BUILDER HELPER
    def add_chapter_to_doc(ch_num, ch_data):
        p_ch = doc.add_paragraph()
        r_chl = p_ch.add_run(f"CHAPTER {ch_num}")
        r_chl.font.bold = True
        r_chl.font.size = Pt(12)
        r_chl.font.color.rgb = RGBColor(37, 99, 235)

        h_ch = doc.add_heading(ch_data["title"].upper(), level=1)
        h_ch.runs[0].font.color.rgb = RGBColor(30, 58, 138)

        for sec_title, sec_content in ch_data["sections"]:
            h_sec = doc.add_heading(sec_title, level=2)
            h_sec.runs[0].font.color.rgb = RGBColor(15, 23, 42)
            for p in sec_content.split('\n\n'):
                if p.startswith("--------------------------------------------------------------------------------"):
                    # Code / algorithm block
                    box = doc.add_table(rows=1, cols=1)
                    box.alignment = WD_TABLE_ALIGNMENT.CENTER
                    cell = box.rows[0].cells[0]
                    set_cell_background(cell, "F8FAFC")
                    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
                    p_box = cell.paragraphs[0]
                    p_box.text = p
                    for run in p_box.runs:
                        run.font.name = "Consolas"
                        run.font.size = Pt(8)
                        run.font.color.rgb = RGBColor(15, 23, 42)
                    doc.add_paragraph()
                else:
                    doc.add_paragraph(p)

    # 8. CHAPTER 1
    add_chapter_to_doc(1, rc.CHAPTER_1)
    doc.add_page_break()

    # 9. CHAPTER 2
    add_chapter_to_doc(2, rc.CHAPTER_2)
    doc.add_paragraph("Table 2.1: Comparative Analysis of Modern Travel and Recommendation Platforms")
    t21 = doc.add_table(rows=len(rc.TABLE_2_1_DATA), cols=len(rc.TABLE_2_1_DATA[0]))
    format_table(t21, rc.TABLE_2_1_DATA, col_widths=[1.5, 1.1, 1.1, 1.1, 1.1, 1.1])
    doc.add_page_break()

    # 10. CHAPTER 3
    add_chapter_to_doc(3, rc.CHAPTER_3)
    doc.add_paragraph("Table 3.1: Functional Requirements Specification (FR-01 to FR-14)")
    t31 = doc.add_table(rows=len(rc.TABLE_3_1_DATA), cols=len(rc.TABLE_3_1_DATA[0]))
    format_table(t31, rc.TABLE_3_1_DATA, col_widths=[0.8, 1.5, 3.8, 0.9])
    doc.add_paragraph()

    doc.add_paragraph("Table 3.2: Non-Functional Requirements Specification")
    t32 = doc.add_table(rows=len(rc.TABLE_3_2_DATA), cols=len(rc.TABLE_3_2_DATA[0]))
    format_table(t32, rc.TABLE_3_2_DATA, col_widths=[0.9, 1.2, 3.2, 1.7])
    doc.add_paragraph()

    doc.add_paragraph("Table 3.3: Hardware and Development Infrastructure Requirements")
    t33 = doc.add_table(rows=len(rc.TABLE_3_3_DATA), cols=len(rc.TABLE_3_3_DATA[0]))
    format_table(t33, rc.TABLE_3_3_DATA, col_widths=[1.8, 2.6, 2.6])
    doc.add_paragraph()

    doc.add_paragraph("Table 3.4: Software and Technology Stack Specifications")
    t34 = doc.add_table(rows=len(rc.TABLE_3_4_DATA), cols=len(rc.TABLE_3_4_DATA[0]))
    format_table(t34, rc.TABLE_3_4_DATA, col_widths=[1.5, 1.3, 0.8, 3.4])
    doc.add_page_break()

    # 11. CHAPTER 4
    add_chapter_to_doc(4, rc.CHAPTER_4)
    add_image_safe(doc, "report_assets/fig_4_1_architecture.png", "Figure 4.1: Three-Tier Client-Server Architecture of TripMeta")
    add_image_safe(doc, "report_assets/fig_4_2_use_case.png", "Figure 4.2: UML Use Case Diagram of the TripMeta Platform")
    add_image_safe(doc, "report_assets/fig_4_3_sequence.png", "Figure 4.3: UML Sequence Diagram: Travel Recommendation Generation Flow")
    add_image_safe(doc, "report_assets/fig_4_4_er_diagram.png", "Figure 4.4: Entity-Relationship (ER) Diagram of SQLite Database")
    
    doc.add_paragraph("Table 4.1: Database Schema: destinations Table")
    t41 = doc.add_table(rows=len(rc.TABLE_4_1_DATA), cols=len(rc.TABLE_4_1_DATA[0]))
    format_table(t41, rc.TABLE_4_1_DATA, col_widths=[1.2, 1.0, 2.0, 2.8])
    doc.add_paragraph()

    doc.add_paragraph("Table 4.2: Database Schema: trips Table")
    t42 = doc.add_table(rows=len(rc.TABLE_4_2_DATA), cols=len(rc.TABLE_4_2_DATA[0]))
    format_table(t42, rc.TABLE_4_2_DATA, col_widths=[1.2, 1.0, 2.0, 2.8])
    doc.add_paragraph()

    doc.add_paragraph("Table 4.3: Database Schema: saved_destinations Table")
    t43 = doc.add_table(rows=len(rc.TABLE_4_3_DATA), cols=len(rc.TABLE_4_3_DATA[0]))
    format_table(t43, rc.TABLE_4_3_DATA, col_widths=[1.2, 1.0, 2.0, 2.8])
    doc.add_paragraph()

    doc.add_paragraph("Table 4.4: RESTful API Endpoints Specification")
    t44 = doc.add_table(rows=len(rc.TABLE_4_4_DATA), cols=len(rc.TABLE_4_4_DATA[0]))
    format_table(t44, rc.TABLE_4_4_DATA, col_widths=[0.9, 1.6, 2.0, 0.9, 1.6])
    doc.add_page_break()

    # 12. CHAPTER 5
    add_chapter_to_doc(5, rc.CHAPTER_5)
    add_image_safe(doc, "report_assets/fig_5_1_algorithm_pipeline.png", "Figure 5.1: TF-IDF Vector Space and Multi-Factor Scoring Pipeline Flowchart")
    doc.add_paragraph("Table 5.1: Computational Complexity of Recommendation Engine Pipeline")
    t51 = doc.add_table(rows=len(rc.TABLE_5_1_DATA), cols=len(rc.TABLE_5_1_DATA[0]))
    format_table(t51, rc.TABLE_5_1_DATA, col_widths=[1.8, 1.2, 1.1, 2.9])
    doc.add_page_break()

    # 13. CHAPTER 6
    add_chapter_to_doc(6, rc.CHAPTER_6)
    doc.add_page_break()

    # 14. CHAPTER 7
    add_chapter_to_doc(7, rc.CHAPTER_7)
    doc.add_paragraph("Table 7.1: Team Member Contributions and Task Allocation Matrix")
    t71 = doc.add_table(rows=len(rc.TABLE_7_1_DATA), cols=len(rc.TABLE_7_1_DATA[0]))
    format_table(t71, rc.TABLE_7_1_DATA, col_widths=[1.8, 1.4, 3.8])
    doc.add_paragraph()

    doc.add_paragraph("Table 7.2: System Integration Challenges and Technical Solutions")
    t72 = doc.add_table(rows=len(rc.TABLE_7_2_DATA), cols=len(rc.TABLE_7_2_DATA[0]))
    format_table(t72, rc.TABLE_7_2_DATA, col_widths=[0.8, 1.7, 2.0, 2.5])
    doc.add_page_break()

    # 15. CHAPTER 8
    add_chapter_to_doc(8, rc.CHAPTER_8)
    doc.add_paragraph("Table 8.1: System Performance and Stress Test Metrics")
    t81 = doc.add_table(rows=len(rc.TABLE_8_1_DATA), cols=len(rc.TABLE_8_1_DATA[0]))
    format_table(t81, rc.TABLE_8_1_DATA, col_widths=[1.1, 0.9, 1.0, 1.1, 1.1, 1.0, 0.8])
    doc.add_paragraph()
    add_image_safe(doc, "report_assets/fig_8_1_performance_benchmark.png", "Figure 8.1: Scalability Benchmarks: Request Latency and Throughput vs Concurrency")
    doc.add_page_break()

    # 16. CHAPTER 9
    add_chapter_to_doc(9, rc.CHAPTER_9)
    add_image_safe(doc, "report_assets/fig_9_1_test_distribution.png", "Figure 9.1: TripMeta Test Suite Coverage and Execution Distribution", max_width=4.5)
    doc.add_paragraph("Table 9.1: Comprehensive System Test Case Execution Matrix (TC-01 to TC-14)")
    t91 = doc.add_table(rows=len(rc.TABLE_9_1_DATA), cols=len(rc.TABLE_9_1_DATA[0]))
    format_table(t91, rc.TABLE_9_1_DATA, col_widths=[0.6, 1.1, 1.0, 1.4, 1.4, 1.1, 0.4])
    doc.add_paragraph()

    doc.add_paragraph("Table 9.2: Validation of Project Objectives Against Experimental Metrics")
    t92 = doc.add_table(rows=len(rc.TABLE_9_2_DATA), cols=len(rc.TABLE_9_2_DATA[0]))
    format_table(t92, rc.TABLE_9_2_DATA, col_widths=[1.5, 1.8, 2.6, 1.1])
    doc.add_page_break()

    # 17. CHAPTER 10
    add_chapter_to_doc(10, rc.CHAPTER_10)
    add_image_safe(doc, "report_assets/fig_10_1_home.png", "Figure 10.1: TripMeta Home Portal: Curated Mood Explorer & Top Highlights", max_width=5.8)
    add_image_safe(doc, "report_assets/fig_10_2_explore.png", "Figure 10.2: Interactive Destination Atlas with State and Category Filters", max_width=5.8)
    add_image_safe(doc, "report_assets/fig_10_3_plan.png", "Figure 10.3: Multi-Criteria Travel Preference Input Wizard ('Plan My Trip')", max_width=5.8)
    add_image_safe(doc, "report_assets/fig_10_4_destination.png", "Figure 10.4: Comprehensive Destination Details, Logistics, and Nearby Spots", max_width=5.8)
    add_image_safe(doc, "report_assets/fig_10_5_saved.png", "Figure 10.5: Saved Destinations and Bookmark Management Dashboard", max_width=5.8)
    add_image_safe(doc, "report_assets/fig_10_6_trip_builder.png", "Figure 10.6: Custom Multi-Day Trip Itinerary Builder with Budget Tracking", max_width=5.8)
    add_image_safe(doc, "report_assets/fig_10_7_results.png", "Figure 10.7: Personalized Recommendation Results with Real-Time Match Scores", max_width=5.8)
    
    doc.add_paragraph("Table 10.1: Feature Comparison: TripMeta vs Commercial OTAs")
    t101 = doc.add_table(rows=len(rc.TABLE_10_1_DATA), cols=len(rc.TABLE_10_1_DATA[0]))
    format_table(t101, rc.TABLE_10_1_DATA, col_widths=[1.5, 1.8, 1.8, 1.9])
    doc.add_page_break()

    # 18. CHAPTER 11
    add_chapter_to_doc(11, rc.CHAPTER_11)
    doc.add_page_break()

    # 19. REFERENCES
    h_ref = doc.add_heading("REFERENCES", level=1)
    h_ref.runs[0].font.color.rgb = RGBColor(30, 58, 138)
    for ref in rc.REFERENCES_LIST:
        p_ref = doc.add_paragraph(ref)
        p_ref.paragraph_format.left_indent = Inches(0.4)
        p_ref.paragraph_format.first_line_indent = Inches(-0.4)
    doc.add_page_break()

    # 20. APPENDIX A
    h_app = doc.add_heading("APPENDIX A", level=1)
    h_app.runs[0].font.color.rgb = RGBColor(30, 58, 138)
    for p in rc.APPENDIX_A_TEXT.split('\n\n'):
        doc.add_paragraph(p)

    doc.save(output_filename)
    print(f"DOCX build complete: {output_filename}")

if __name__ == "__main__":
    create_docx()
